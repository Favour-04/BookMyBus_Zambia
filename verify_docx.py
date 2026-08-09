"""Verify the generated Word document follows SICT Project Guidelines."""
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

doc = Document('GROUP_29_PROJECT_REPORT_CORRECTED.docx')

print(f'Total paragraphs: {len(doc.paragraphs)}')
print(f'Total tables: {len(doc.tables)}')
print(f'Total sections: {len(doc.sections)}')
print()

# ── 1. Check default font ──
style = doc.styles['Normal']
print('=== FONT CHECK ===')
print(f'Default font name: {style.font.name}')
print(f'Default font size: {style.font.size}')
if style.font.name == 'Times New Roman' and style.font.size == Pt(12):
    print('✅ Font: Times New Roman 12pt')
else:
    print('❌ Font: Expected Times New Roman 12pt')
print()

# ── 2. Check line spacing ──
pf = style.paragraph_format
print('=== LINE SPACING CHECK ===')
print(f'Default line spacing: {pf.line_spacing}')
if pf.line_spacing == 1.5:
    print('✅ Line spacing: 1.5')
else:
    print('❌ Line spacing: Expected 1.5')
print()

# ── 3. Check margins ──
print('=== MARGINS CHECK ===')
for i, section in enumerate(doc.sections):
    print(f'Section {i}: top={section.top_margin}, bottom={section.bottom_margin}, '
          f'left={section.left_margin}, right={section.right_margin}')
    if (section.top_margin == Inches(1) and section.bottom_margin == Inches(1) and
            section.left_margin == Inches(1) and section.right_margin == Inches(1)):
        print(f'  ✅ Section {i} margins: 1 inch all sides')
    else:
        print(f'  ❌ Section {i} margins: Expected 1 inch all sides')
print()

# ── 4. Check page numbering ──
print('=== PAGE NUMBERING CHECK ===')
for i, section in enumerate(doc.sections):
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    fmt = pgNumType.get(qn('w:fmt')) if pgNumType is not None else 'not set'
    start = pgNumType.get(qn('w:start')) if pgNumType is not None else 'not set'
    diff_first = section.different_first_page_header_footer
    print(f'Section {i}: page format={fmt}, start={start}, different_first={diff_first}')
    # Check footer
    footer = section.footer
    footer_text = ' '.join(p.text for p in footer.paragraphs)
    has_page_field = 'PAGE' in footer_text or any(
        'PAGE' in p._p.xml for p in footer.paragraphs
    )
    print(f'  Footer has page number field: {has_page_field}')
print()

# ── 5. Check heading styles ──
print('=== HEADING STYLES CHECK ===')
for level in range(1, 5):
    heading_style = doc.styles[f'Heading {level}']
    print(f'Heading {level}: font={heading_style.font.name}, '
          f'size={heading_style.font.size}, bold={heading_style.font.bold}')
print()

# ── 6. Check first 8 paragraphs ──
print('=== FIRST 8 PARAGRAPHS ===')
for p in doc.paragraphs[:8]:
    print(f'  [{p.style.name}] {p.text[:100]}')
print()

# ── 7. Check heading distribution ──
headings = {}
for p in doc.paragraphs:
    if p.style.name.startswith('Heading'):
        level = p.style.name
        headings[level] = headings.get(level, 0) + 1

print('=== HEADING COUNTS ===')
print(headings)
print()

# ── 8. Check tables ──
print('=== TABLES ===')
print(f'Tables found: {len(doc.tables)}')
if doc.tables:
    print(f'First table dimensions: {len(doc.tables[0].rows)} rows x {len(doc.tables[0].columns)} cols')
    print(f'First table header: {[c.text for c in doc.tables[0].rows[0].cells]}')
print()

# ── 9. Check abstract single spacing ──
print('=== ABSTRACT SPACING CHECK ===')
in_abstract = False
abstract_spacing_ok = True
for p in doc.paragraphs:
    if p.style.name == 'Heading 2' and p.text.strip() == 'ABSTRACT':
        in_abstract = True
        continue
    if in_abstract and p.style.name.startswith('Heading'):
        break
    if in_abstract and p.text.strip():
        ls = p.paragraph_format.line_spacing
        if ls is not None and ls != 1.0:
            abstract_spacing_ok = False
            print(f'  ❌ Abstract paragraph not single-spaced: "{p.text[:50]}" (spacing={ls})')
if abstract_spacing_ok:
    print('✅ Abstract is single-spaced')
print()

print('=== VERIFICATION COMPLETE ===')