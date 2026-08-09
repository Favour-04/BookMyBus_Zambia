"""
Convert GROUP_29_PROJECT_REPORT_CORRECTED.md to a Word document (.docx)
following SICT Project Guidelines:
- Times New Roman 12pt body text
- 1.5 line spacing (single for abstract, code blocks, references)
- 1-inch margins on all sides
- Roman numeral page numbering for front matter (title page counted but not numbered)
- Arabic numeral page numbering (bottom-right) for main body
"""

import re
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement

INPUT_FILE = "GROUP_29_PROJECT_REPORT_CORRECTED.md"
OUTPUT_FILE = "GROUP_29_PROJECT_REPORT_CORRECTED.docx"

# ── Helpers ──────────────────────────────────────────────────────────────

def set_run_font(run, name='Times New Roman', size=Pt(12), bold=None, italic=None):
    """Set font properties on a run, including East Asian font."""
    run.font.name = name
    run.font.size = size
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:eastAsia'), name)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_page_number_format(section, fmt='decimal', start_at=None):
    """Set the page number format for a section."""
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    pgNumType.set(qn('w:fmt'), fmt)
    if start_at is not None:
        pgNumType.set(qn('w:start'), str(start_at))


def add_page_number_field(paragraph):
    """Add a PAGE field to a paragraph."""
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    set_run_font(run, size=Pt(12))
    return run


def setup_footer(section, fmt='decimal', start_at=None, different_first=False):
    """Set up footer with page number at bottom-right."""
    section.different_first_page_header_footer = different_first
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    # Clear existing runs
    for run in p.runs:
        run.text = ''
    add_page_number_field(p)
    set_page_number_format(section, fmt, start_at)


def add_inline_formatting(paragraph, text, font_size=Pt(12)):
    """Parse **bold**, *italic*, and `code` inline markers and add runs."""
    pattern = re.compile(r'(\*\*.+?\*\*|\*.+?\*|`[^`]+`)')
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            set_run_font(run, size=font_size)
        elif part.startswith('*') and part.endswith('*') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            set_run_font(run, size=font_size)
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name='Consolas', size=Pt(10))
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=font_size)


def set_cell_shading(cell, color_hex):
    """Apply background shading to a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def add_table(doc, rows):
    """Add a Markdown table to the document with header styling."""
    if not rows:
        return

    num_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            if j >= num_cols:
                break
            cell = table.cell(i, j)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            add_inline_formatting(p, cell_text.strip(), font_size=Pt(10))
            if i == 0:
                set_cell_shading(cell, 'D9E2F3')
                for run in p.runs:
                    run.bold = True

    # Add spacing after table
    doc.add_paragraph()


def add_code_block(doc, code_lines):
    """Add a code block as monospaced paragraphs with light shading."""
    for line in code_lines:
        p = doc.add_paragraph()
        run = p.add_run(line if line else ' ')
        set_run_font(run, name='Consolas', size=Pt(9))
        pPr = p._p.get_or_add_pPr()
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F2F2F2"/>')
        pPr.append(shading)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
    # Add spacing after code block
    doc.add_paragraph()


def add_list_item(doc, text, numbered=False, level=0):
    """Add a list item paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5 + 0.25 * level)
    prefix = '• '
    if prefix:
        run = p.add_run(prefix)
        run.bold = True
        set_run_font(run)
    add_inline_formatting(p, text)
    return p


# ── Main conversion ──────────────────────────────────────────────────────

def convert():
    with open(INPUT_FILE, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # ── Set default font: Times New Roman 12pt ──
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    # ── Default paragraph format: 1.5 line spacing ──
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)

    # ── Margins: 1 inch on all sides ──
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── Configure heading styles: Times New Roman, black ──
    heading_configs = {
        1: {'size': Pt(16), 'bold': True, 'italic': False},
        2: {'size': Pt(14), 'bold': True, 'italic': False},
        3: {'size': Pt(12), 'bold': True, 'italic': False},
        4: {'size': Pt(12), 'bold': True, 'italic': True},
    }
    for level, config in heading_configs.items():
        heading_style = doc.styles[f'Heading {level}']
        heading_style.font.name = 'Times New Roman'
        heading_style.font.size = config['size']
        heading_style.font.bold = config['bold']
        heading_style.font.italic = config['italic']
        heading_style.font.color.rgb = RGBColor(0, 0, 0)
        h_rPr = heading_style.element.get_or_add_rPr()
        h_rFonts = h_rPr.find(qn('w:rFonts'))
        if h_rFonts is None:
            h_rFonts = OxmlElement('w:rFonts')
            h_rPr.append(h_rFonts)
        h_rFonts.set(qn('w:ascii'), 'Times New Roman')
        h_rFonts.set(qn('w:hAnsi'), 'Times New Roman')
        h_rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    # ── State tracking ──
    in_front_matter = True
    in_abstract = False
    in_references = False
    main_section = None

    i = 0
    n = len(lines)
    in_code_block = False
    code_buffer = []

    while i < n:
        line = lines[i].rstrip('\n')

        # Code block handling
        if line.strip().startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_buffer = []
            else:
                in_code_block = False
                add_code_block(doc, code_buffer)
            i += 1
            continue

        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue

        stripped = line.strip()

        # Skip empty lines
        if not stripped:
            i += 1
            continue

        # ── Detect transition to main body (Chapter 1) ──
        if in_front_matter and stripped.startswith('# CHAPTER 1'):
            main_section = doc.add_section(WD_SECTION.NEW_PAGE)
            main_section.top_margin = Inches(1)
            main_section.bottom_margin = Inches(1)
            main_section.left_margin = Inches(1)
            main_section.right_margin = Inches(1)
            # Front matter: Roman numerals, title page not numbered
            setup_footer(doc.sections[0], fmt='upperRoman', different_first=True)
            # Main body: Arabic numerals starting at 1
            setup_footer(main_section, fmt='decimal', start_at=1)
            in_front_matter = False

        # ── Track abstract section for single spacing ──
        if stripped.startswith('## ABSTRACT'):
            in_abstract = True
        elif in_abstract and stripped.startswith('## '):
            in_abstract = False

        # ── Track references section for single spacing ──
        if stripped.startswith('## REFERENCES'):
            in_references = True
        elif in_references and stripped.startswith('# '):
            in_references = False

        # Horizontal rule → page break
        if re.match(r'^-{3,}$', stripped):
            doc.add_page_break()
            i += 1
            continue

        # Headings
        heading_match = re.match(r'^(#{1,4})\s+(.*)', stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            text = re.sub(r'\*(.+?)\*', r'\1', text)
            text = re.sub(r'`(.+?)`', r'\1', text)
            doc.add_heading(text, level=level)
            i += 1
            continue

        # Table detection
        if stripped.startswith('|') and stripped.endswith('|'):
            table_rows = []
            while i < n:
                tline = lines[i].rstrip('\n').strip()
                if not (tline.startswith('|') and tline.endswith('|')):
                    break
                if re.match(r'^\|[\s\-:|]+\|$', tline):
                    i += 1
                    continue
                cells = [c.strip() for c in tline.strip('|').split('|')]
                table_rows.append(cells)
                i += 1
            add_table(doc, table_rows)
            continue

        # Unordered list
        if re.match(r'^[-*]\s+', stripped):
            text = re.sub(r'^[-*]\s+', '', stripped)
            add_list_item(doc, text, numbered=False)
            i += 1
            continue

        # Ordered list
        ordered_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if ordered_match:
            num = ordered_match.group(1)
            text = ordered_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            run = p.add_run(f'{num}. ')
            run.bold = True
            set_run_font(run)
            add_inline_formatting(p, text)
            i += 1
            continue

        # Signature lines (lines of underscores)
        if re.match(r'^_{5,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(stripped)
            run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            set_run_font(run)
            i += 1
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        # Single spacing for abstract and references per guidelines
        if in_abstract or in_references:
            p.paragraph_format.line_spacing = 1.0
        add_inline_formatting(p, stripped)
        i += 1

    # Handle any unclosed code block
    if in_code_block:
        add_code_block(doc, code_buffer)

    # If we never hit Chapter 1, set up footer for the single section
    if in_front_matter:
        setup_footer(doc.sections[0], fmt='upperRoman', different_first=True)

    doc.save(OUTPUT_FILE)
    print(f"Successfully converted to {OUTPUT_FILE}")


if __name__ == '__main__':
    convert()